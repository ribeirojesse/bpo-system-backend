"""Exportação em Excel (.xlsx) dos lançamentos do relatório.

Uma planilha "Lançamentos" com um lançamento por linha (todas as
informações dele: tipo, datas, competência, contato financeiro e
documento, descrição, categoria, subcategoria, valor, conta bancária,
conciliação e observação) — ver DreService.listar_lancamentos.

O único filtro é o período personalizado (data inicial e final), que é
o que aparece no cabeçalho da planilha.

No topo, um resumo com Receitas / Despesas / Resultado calculados por
fórmula (SOMASE) sobre a própria tabela — se a pessoa editar ou apagar
linhas no Excel, o resumo acompanha. Datas e valores são gravados como
data e número de verdade (não texto), então filtros, ordenação, somas e
tabelas dinâmicas funcionam direto.
"""

import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# (título, chave, largura, formato)
_COLUNAS = [
    ("Tipo", "tipo", 10, None),
    ("Data", "data", 12, "DD/MM/YYYY"),
    ("Vencimento", "vencimento", 12, "DD/MM/YYYY"),
    ("Competência", "competencia", 12, None),
    ("Contato financeiro", "contato", 32, None),
    ("CPF/CNPJ do contato", "contato_documento", 20, None),
    ("Descrição", "descricao", 42, None),
    ("Categoria", "categoria", 24, None),
    ("Subcategoria", "subcategoria", 24, None),
    ("Valor", "valor", 15, '"R$" #,##0.00'),
    ("Conta bancária", "conta_bancaria", 28, None),
    ("Conciliado", "conciliado", 11, None),
    ("Observação", "observacao", 40, None),
]

_FORMATO_MOEDA = '"R$" #,##0.00'

# Paleta "Notion" usada também no PDF do DRE (dre_pdf_service.py).
_COR_TEXTO = "37352F"
_COR_SECUNDARIA = "9B9A97"
_COR_BORDA = "E9E9E7"
_COR_ZEBRA = "F7F6F5"
_COR_RECEITA = "15803D"
_COR_DESPESA = "B91C1C"

_LINHA_CABECALHO = 8


def nome_arquivo(cliente_nome, data_inicio, data_fim):
    """Nome de arquivo só com ASCII (cabeçalho HTTP seguro)."""

    base = unicodedata.normalize("NFKD", cliente_nome or "cliente")
    base = base.encode("ascii", "ignore").decode("ascii")
    base = re.sub(r"[^A-Za-z0-9]+", "_", base).strip("_") or "cliente"

    return (
        f"Lancamentos_{base[:60]}_"
        f"{data_inicio:%d-%m-%Y}_a_{data_fim:%d-%m-%Y}.xlsx"
    )


def gerar(dados) -> bytes:

    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"

    lancamentos = dados["lancamentos"]

    # ---------------- título ----------------
    ws["A1"] = f"Lançamentos — {dados['cliente_nome']}"
    ws["A1"].font = Font(bold=True, size=14, color=_COR_TEXTO)

    ws["A2"] = (
        f"Período: {dados['data_inicio']:%d/%m/%Y} "
        f"a {dados['data_fim']:%d/%m/%Y}"
    )
    ws["A2"].font = Font(size=11, color=_COR_TEXTO)

    # ---------------- resumo (fórmulas) ----------------
    primeira = _LINHA_CABECALHO + 1
    ultima = _LINHA_CABECALHO + max(len(lancamentos), 1)

    col_tipo = "A"
    col_valor = get_column_letter(
        [c[1] for c in _COLUNAS].index("valor") + 1
    )

    faixa_tipo = f"${col_tipo}${primeira}:${col_tipo}${ultima}"
    faixa_valor = f"${col_valor}${primeira}:${col_valor}${ultima}"

    resumo = [
        (4, "Receitas", f'=SUMIF({faixa_tipo},"Receita",{faixa_valor})',
         _COR_RECEITA),
        (5, "Despesas", f'=SUMIF({faixa_tipo},"Despesa",{faixa_valor})',
         _COR_DESPESA),
        (6, "Resultado", "=B4-B5", _COR_TEXTO),
    ]

    for linha, rotulo, formula, cor in resumo:
        ws.cell(row=linha, column=1, value=rotulo).font = Font(
            bold=True, color=_COR_TEXTO
        )
        celula = ws.cell(row=linha, column=2, value=formula)
        celula.number_format = _FORMATO_MOEDA
        celula.font = Font(bold=True, color=cor)

    ws.cell(row=4, column=3, value=f"{len(lancamentos)} lançamento(s)").font = (
        Font(size=10, color=_COR_SECUNDARIA)
    )

    # ---------------- cabeçalho da tabela ----------------
    borda = Border(bottom=Side(style="thin", color=_COR_BORDA))
    fundo_cabecalho = PatternFill("solid", fgColor=_COR_TEXTO)
    fundo_zebra = PatternFill("solid", fgColor=_COR_ZEBRA)

    for idx, (titulo, _chave, largura, _fmt) in enumerate(_COLUNAS, start=1):

        celula = ws.cell(row=_LINHA_CABECALHO, column=idx, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = fundo_cabecalho
        celula.alignment = Alignment(vertical="center")

        ws.column_dimensions[get_column_letter(idx)].width = largura

    ws.row_dimensions[_LINHA_CABECALHO].height = 20

    # ---------------- linhas ----------------
    if not lancamentos:

        ws.cell(
            row=primeira,
            column=1,
            value="Nenhum lançamento no período."
        ).font = Font(italic=True, color=_COR_SECUNDARIA)

    for n, lanc in enumerate(lancamentos):

        linha = primeira + n

        for idx, (_titulo, chave, _largura, fmt) in enumerate(
            _COLUNAS, start=1
        ):
            valor = lanc.get(chave)

            if chave == "valor":
                valor = float(valor or 0)
            elif chave == "conciliado":
                valor = "Sim" if valor else "Não"
            elif valor is None:
                valor = ""

            celula = ws.cell(row=linha, column=idx, value=valor)

            if fmt:
                celula.number_format = fmt

            celula.border = borda

            if n % 2 == 1:
                celula.fill = fundo_zebra

            if chave == "valor":
                celula.font = Font(
                    color=(
                        _COR_RECEITA
                        if lanc["tipo"] == "Receita"
                        else _COR_DESPESA
                    )
                )

            if chave in ("descricao", "observacao"):
                celula.alignment = Alignment(wrap_text=False)

    # ---------------- congelar cabeçalho ----------------
    # Sem autofiltro nas colunas: o único filtro da planilha é o período
    # (pedido do usuário). Quem quiser filtrar liga no Excel (Ctrl+Shift+L).
    ws.freeze_panes = f"A{primeira}"

    ws.sheet_view.showGridLines = False

    # Impressão: paisagem, cabeçalho repetido em cada página.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{_LINHA_CABECALHO}:{_LINHA_CABECALHO}"

    buffer = io.BytesIO()
    wb.save(buffer)

    return buffer.getvalue()
